"""
Funciones para generar cotizaciones en Word.
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt


CARPETA_DOCUMENTOS = Path("documentos")


def generar_cotizacion(
    nombre_cliente,
    nit,
    direccion,
    producto,
    cantidad,
    precio
):
    """
    Genera una cotización en formato Word.
    """

    # Crear la carpeta documentos si no existe
    CARPETA_DOCUMENTOS.mkdir(
        exist_ok=True
    )

    # Calcular el total
    total = cantidad * precio

    # Crear documento Word
    documento = Document()

    # Título
    titulo = documento.add_paragraph()

    texto_titulo = titulo.add_run(
        "COTIZACIÓN"
    )

    texto_titulo.bold = True
    texto_titulo.font.size = Pt(20)

    titulo.alignment = 1


    # Datos del cliente
    documento.add_paragraph(
        f"Cliente: {nombre_cliente}"
    )

    documento.add_paragraph(
        f"NIT: {nit}"
    )

    documento.add_paragraph(
        f"Dirección: {direccion}"
    )


    # Espacio
    documento.add_paragraph()


    # Tabla
    tabla = documento.add_table(
        rows=1,
        cols=4
    )

    tabla.style = "Table Grid"


    encabezados = [
        "Producto",
        "Cantidad",
        "Precio",
        "Total"
    ]


    for columna, encabezado in enumerate(encabezados):

        tabla.rows[0].cells[columna].text = encabezado


    # Agregar producto
    fila = tabla.add_row().cells

    fila[0].text = str(producto)
    fila[1].text = str(cantidad)
    fila[2].text = f"Q{precio:.2f}"
    fila[3].text = f"Q{total:.2f}"


    # Total general
    documento.add_paragraph()

    total_parrafo = documento.add_paragraph()

    texto_total = total_parrafo.add_run(
        f"TOTAL: Q{total:.2f}"
    )

    texto_total.bold = True
    texto_total.font.size = Pt(14)


    # Nombre del archivo
    nombre_archivo = (
        f"Cotizacion_{nombre_cliente}.docx"
    )

    # Evitar problemas con algunos caracteres
    nombre_archivo = nombre_archivo.replace(
        "/",
        "_"
    )

    nombre_archivo = nombre_archivo.replace(
        "\\",
        "_"
    )


    ruta_archivo = (
        CARPETA_DOCUMENTOS /
        nombre_archivo
    )


    # Guardar documento
    documento.save(
        ruta_archivo
    )


    return ruta_archivo