"""
Funciones para generar informes del proyecto.
"""

from openpyxl import load_workbook
from pathlib import Path


ARCHIVO_EXCEL = Path("datos") / "Inventario.xlsx"


def obtener_informe_productos():
    """
    Obtiene todos los productos registrados.
    """

    libro = load_workbook(ARCHIVO_EXCEL)
    hoja = libro["Productos"]

    productos = []

    for fila in hoja.iter_rows(min_row=2, values_only=True):

        if fila[0] is not None:

            producto = {
                "nombre": fila[0],
                "precio": fila[1],
                "existencia": fila[2]
            }

            productos.append(producto)

    libro.close()

    return productos


def obtener_informe_clientes():
    """
    Obtiene todos los clientes registrados.
    """

    libro = load_workbook(ARCHIVO_EXCEL)
    hoja = libro["Clientes"]

    clientes = []

    for fila in hoja.iter_rows(min_row=2, values_only=True):

        if fila[0] is not None:

            cliente = {
                "nombre": fila[0],
                "nit": fila[1],
                "direccion": fila[2]
            }

            clientes.append(cliente)

    libro.close()

    return clientes


def obtener_informe_pedidos():
    """
    Obtiene todos los pedidos registrados.
    """

    libro = load_workbook(ARCHIVO_EXCEL)
    hoja = libro["Pedidos"]

    pedidos = []

    for fila in hoja.iter_rows(min_row=2, values_only=True):

        if fila[0] is not None:

            pedido = {
                "cliente": fila[0],
                "producto": fila[1],
                "cantidad": fila[2],
                "valor": fila[3]
            }

            pedidos.append(pedido)

    libro.close()

    return pedidos


def calcular_total_ventas():
    """
    Calcula el valor total de todos los pedidos.
    """

    pedidos = obtener_informe_pedidos()

    total = 0

    for pedido in pedidos:

        total += pedido["valor"]

    return total


def contar_productos():
    """
    Cuenta la cantidad de productos registrados.
    """

    productos = obtener_informe_productos()

    return len(productos)


def contar_clientes():
    """
    Cuenta la cantidad de clientes registrados.
    """

    clientes = obtener_informe_clientes()

    return len(clientes)


def contar_pedidos():
    """
    Cuenta la cantidad de pedidos registrados.
    """

    pedidos = obtener_informe_pedidos()

    return len(pedidos)