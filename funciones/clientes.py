"""
Funciones para administrar los clientes.
"""

from openpyxl import load_workbook
from pathlib import Path


# ==========================================================
# UBICACIÓN DEL ARCHIVO EXCEL
# ==========================================================

ARCHIVO_EXCEL = Path("datos") / "Inventario.xlsx"


# ==========================================================
# AGREGAR CLIENTE
# ==========================================================

def agregar_cliente(nombre, nit, direccion):
    """
    Agrega un nuevo cliente al archivo Excel.
    """

    libro = load_workbook(ARCHIVO_EXCEL)

    hoja = libro["Clientes"]

    hoja.append([
        nombre,
        nit,
        direccion
    ])

    libro.save(ARCHIVO_EXCEL)
    libro.close()


# ==========================================================
# LISTAR CLIENTES
# ==========================================================

def listar_clientes():
    """
    Devuelve todos los clientes registrados.
    """

    libro = load_workbook(ARCHIVO_EXCEL)

    hoja = libro["Clientes"]

    clientes = []

    for fila in hoja.iter_rows(min_row=2, values_only=True):

        if fila[0] is not None:

            cliente = {
                "nombre": fila[0],
                "nit": fila[1],
                "direccion": fila[2]
            }

            clientes.append(cliente)

    libro.close()

    return clientes


# ==========================================================
# BUSCAR CLIENTE
# ==========================================================

def buscar_cliente(nombre):
    """
    Busca un cliente por su nombre.
    """

    clientes = listar_clientes()

    for cliente in clientes:

        if cliente["nombre"].lower() == nombre.lower():
            return cliente

    return None


# ==========================================================
# EDITAR CLIENTE
# ==========================================================

def editar_cliente(
    nombre_actual,
    nuevo_nombre,
    nuevo_nit,
    nueva_direccion
):
    """
    Edita un cliente existente.
    """

    libro = load_workbook(ARCHIVO_EXCEL)

    hoja = libro["Clientes"]

    encontrado = False

    for fila in range(2, hoja.max_row + 1):

        nombre = hoja.cell(fila, 1).value

        if (
            nombre is not None
            and nombre.lower() == nombre_actual.lower()
        ):

            hoja.cell(fila, 1).value = nuevo_nombre
            hoja.cell(fila, 2).value = nuevo_nit
            hoja.cell(fila, 3).value = nueva_direccion

            encontrado = True
            break

    libro.save(ARCHIVO_EXCEL)
    libro.close()

    return encontrado


# ==========================================================
# ELIMINAR CLIENTE
# ==========================================================

def eliminar_cliente(nombre):
    """
    Elimina un cliente por su nombre.
    """

    libro = load_workbook(ARCHIVO_EXCEL)

    hoja = libro["Clientes"]

    encontrado = False

    for fila in range(2, hoja.max_row + 1):

        nombre_actual = hoja.cell(fila, 1).value

        if (
            nombre_actual is not None
            and nombre_actual.lower() == nombre.lower()
        ):

            hoja.delete_rows(fila, 1)

            encontrado = True
            break

    libro.save(ARCHIVO_EXCEL)
    libro.close()

    return encontrado