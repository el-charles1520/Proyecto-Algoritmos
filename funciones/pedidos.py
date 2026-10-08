"""
Funciones para administrar los pedidos.
"""

from openpyxl import load_workbook
from pathlib import Path


# ==========================================================
# UBICACIÓN DEL ARCHIVO EXCEL
# ==========================================================

ARCHIVO_EXCEL = Path("datos") / "Inventario.xlsx"


# ==========================================================
# AGREGAR PEDIDO
# ==========================================================

def agregar_pedido(cliente, producto, cantidad):
    """
    Agrega un pedido y calcula automáticamente su valor.

    El valor se calcula utilizando:
    
    Valor = Precio del producto × Cantidad
    """

    libro = load_workbook(ARCHIVO_EXCEL)

    hoja_productos = libro["Productos"]
    hoja_pedidos = libro["Pedidos"]

    precio_producto = None

    # Buscar el producto
    for fila in range(2, hoja_productos.max_row + 1):

        nombre_producto = hoja_productos.cell(
            fila,
            1
        ).value

        if (
            nombre_producto is not None
            and nombre_producto.lower() == producto.lower()
        ):

            precio_producto = hoja_productos.cell(
                fila,
                2
            ).value

            break

    # Verificar si el producto existe
    if precio_producto is None:

        libro.close()

        return False

    # Calcular valor del pedido
    valor = precio_producto * cantidad

    # Registrar pedido
    hoja_pedidos.append([
        cliente,
        producto,
        cantidad,
        valor
    ])

    libro.save(ARCHIVO_EXCEL)

    libro.close()

    return True


# ==========================================================
# LISTAR PEDIDOS
# ==========================================================

def listar_pedidos():
    """
    Devuelve todos los pedidos registrados.
    """

    libro = load_workbook(ARCHIVO_EXCEL)

    hoja = libro["Pedidos"]

    pedidos = []

    for fila in hoja.iter_rows(
        min_row=2,
        values_only=True
    ):

        if fila[0] is not None:

            pedido = {
                "cliente": fila[0],
                "producto": fila[1],
                "cantidad": fila[2],
                "valor": fila[3]
            }

            pedidos.append(pedido)

    libro.close()

    return pedidos


# ==========================================================
# BUSCAR PEDIDOS DE UN CLIENTE
# ==========================================================

def buscar_pedidos_cliente(cliente):
    """
    Busca todos los pedidos realizados
    por un cliente.
    """

    pedidos = listar_pedidos()

    resultados = []

    for pedido in pedidos:

        if pedido["cliente"].lower() == cliente.lower():

            resultados.append(pedido)

    return resultados


# ==========================================================
# ELIMINAR PEDIDO
# ==========================================================

def eliminar_pedido(cliente, producto):
    """
    Elimina el primer pedido encontrado
    de un cliente y producto.
    """

    libro = load_workbook(ARCHIVO_EXCEL)

    hoja = libro["Pedidos"]

    encontrado = False

    for fila in range(
        2,
        hoja.max_row + 1
    ):

        cliente_actual = hoja.cell(
            fila,
            1
        ).value

        producto_actual = hoja.cell(
            fila,
            2
        ).value

        if (
            cliente_actual is not None
            and producto_actual is not None
            and cliente_actual.lower() == cliente.lower()
            and producto_actual.lower() == producto.lower()
        ):

            hoja.delete_rows(
                fila,
                1
            )

            encontrado = True

            break

    libro.save(ARCHIVO_EXCEL)

    libro.close()

    return encontrado