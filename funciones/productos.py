"""
Funciones para administrar los productos.
"""

from openpyxl import load_workbook
from pathlib import Path


# ==========================================================
# UBICACIÓN DEL ARCHIVO EXCEL
# ==========================================================

ARCHIVO_EXCEL = Path("datos") / "Inventario.xlsx"


# ==========================================================
# AGREGAR PRODUCTO
# ==========================================================

def agregar_producto(nombre, precio, existencia):
    """
    Agrega un nuevo producto al archivo Excel.
    """

    libro = load_workbook(ARCHIVO_EXCEL)

    hoja = libro["Productos"]

    hoja.append([
        nombre,
        precio,
        existencia
    ])

    libro.save(ARCHIVO_EXCEL)
    libro.close()


# ==========================================================
# LISTAR PRODUCTOS
# ==========================================================

def listar_productos():
    """
    Devuelve todos los productos registrados.
    """

    libro = load_workbook(ARCHIVO_EXCEL)

    hoja = libro["Productos"]

    productos = []

    for fila in hoja.iter_rows(min_row=2, values_only=True):

        if fila[0] is not None:

            producto = {
                "nombre": fila[0],
                "precio": fila[1],
                "existencia": fila[2]
            }

            productos.append(producto)

    libro.close()

    return productos


# ==========================================================
# BUSCAR PRODUCTO
# ==========================================================

def buscar_producto(nombre):
    """
    Busca un producto por su nombre.
    """

    productos = listar_productos()

    for producto in productos:

        if producto["nombre"].lower() == nombre.lower():
            return producto

    return None


# ==========================================================
# EDITAR PRODUCTO
# ==========================================================

def editar_producto(
    nombre_actual,
    nuevo_nombre,
    nuevo_precio,
    nueva_existencia
):
    """
    Edita un producto existente.
    """

    libro = load_workbook(ARCHIVO_EXCEL)

    hoja = libro["Productos"]

    encontrado = False

    for fila in range(2, hoja.max_row + 1):

        nombre = hoja.cell(fila, 1).value

        if (
            nombre is not None
            and nombre.lower() == nombre_actual.lower()
        ):

            hoja.cell(fila, 1).value = nuevo_nombre
            hoja.cell(fila, 2).value = nuevo_precio
            hoja.cell(fila, 3).value = nueva_existencia

            encontrado = True
            break

    libro.save(ARCHIVO_EXCEL)
    libro.close()

    return encontrado


# ==========================================================
# ELIMINAR PRODUCTO
# ==========================================================

def eliminar_producto(nombre):
    """
    Elimina un producto por su nombre.
    """

    libro = load_workbook(ARCHIVO_EXCEL)

    hoja = libro["Productos"]

    encontrado = False

    for fila in range(2, hoja.max_row + 1):

        nombre_actual = hoja.cell(fila, 1).value

        if (
            nombre_actual is not None
            and nombre_actual.lower() == nombre.lower()
        ):

            hoja.delete_rows(fila, 1)

            encontrado = True
            break

    libro.save(ARCHIVO_EXCEL)
    libro.close()

    return encontrado